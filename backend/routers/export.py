"""
Export endpoints — convert query results and chat sessions to Excel.
"""
import io
import json
from datetime import datetime
from typing import Any, Dict, List, Optional

import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

router = APIRouter(prefix="/api/export", tags=["Export"])


class ExportRequest(BaseModel):
    filename: Optional[str] = None
    sheets: List[Dict[str, Any]]
    # Each sheet: { "name": str, "columns": list[str], "rows": list[list] }


class QueryExportRequest(BaseModel):
    filename: Optional[str] = None
    title: Optional[str] = None
    columns: List[str]
    rows: List[List[Any]]
    sql: Optional[str] = None
    metadata: Optional[Dict[str, Any]] = None


def _build_excel_response(buffer: io.BytesIO, filename: str) -> StreamingResponse:
    buffer.seek(0)
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    return StreamingResponse(buffer, headers=headers)


@router.post("/excel/query")
def export_query_to_excel(request: QueryExportRequest):
    """Export a single query result to a formatted Excel file."""
    filename = request.filename or f"export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Data sheet
        df = pd.DataFrame(request.rows, columns=request.columns)
        df.to_excel(writer, sheet_name="Data", index=False)

        # Style the data sheet
        ws = writer.sheets["Data"]
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for col_idx, col_name in enumerate(request.columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

            # Auto-width
            max_len = max(
                len(str(col_name)),
                *[len(str(row[col_idx - 1])) for row in request.rows if request.rows],
                10,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

        # Metadata sheet
        meta_data = {
            "Field": ["Export Date", "Row Count", "Title"],
            "Value": [
                datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
                len(request.rows),
                request.title or "Query Export",
            ],
        }
        if request.sql:
            meta_data["Field"].append("SQL Query")
            meta_data["Value"].append(request.sql)
        if request.metadata:
            for k, v in request.metadata.items():
                meta_data["Field"].append(k)
                meta_data["Value"].append(str(v))

        pd.DataFrame(meta_data).to_excel(writer, sheet_name="Metadata", index=False)

    return _build_excel_response(buffer, filename)


@router.post("/excel/multi-sheet")
def export_multi_sheet(request: ExportRequest):
    """Export multiple datasets as separate Excel sheets."""
    filename = request.filename or f"export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet in request.sheets:
            name = sheet.get("name", "Sheet")[:31]  # Excel sheet name limit
            columns = sheet.get("columns", [])
            rows = sheet.get("rows", [])
            df = pd.DataFrame(rows, columns=columns)
            df.to_excel(writer, sheet_name=name, index=False)

    return _build_excel_response(buffer, filename)


@router.post("/excel/session/{session_id}")
def export_session_to_excel(session_id: str):
    """Export a full chat session (messages + any embedded query results) to Excel."""
    from backend.database import db, COLL_MESSAGES, COLL_SESSIONS

    session = db.get(COLL_SESSIONS, session_id)
    if not session:
        raise HTTPException(404, "Session not found")

    messages = db.get_list(COLL_MESSAGES, session_id)
    if not messages:
        raise HTTPException(404, "No messages in session")

    filename = f"session_{session_id[:8]}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Chat history sheet
        chat_rows = []
        for msg in messages:
            chat_rows.append({
                "Timestamp": msg.get("timestamp", ""),
                "Role": msg.get("role", ""),
                "Content": msg.get("content", ""),
            })
        pd.DataFrame(chat_rows).to_excel(writer, sheet_name="Chat History", index=False)

        # Extract any embedded query results from assistant messages
        sheet_idx = 1
        for msg in messages:
            meta = msg.get("metadata", {})
            if meta.get("query_result") and meta["query_result"].get("success"):
                qr = meta["query_result"]
                df = pd.DataFrame(qr.get("rows", []), columns=qr.get("columns", []))
                sheet_name = f"Query_{sheet_idx}"[:31]
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                sheet_idx += 1

    return _build_excel_response(buffer, filename)
